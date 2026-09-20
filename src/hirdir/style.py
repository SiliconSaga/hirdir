"""Shared fonts, fills and small openpyxl helpers for every sheet."""

from __future__ import annotations

from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

FONT = "Arial"
GREEN = "1F5C3A"  # headers
LINK = "008000"  # values pulled from another sheet
INPUT_TEXT = "0000FF"  # cells a human types into
PALE_RED = "F8D0CC"
PALE_GREEN = "C8E6C9"
GREY = "DDDDDD"
LIGHT = "EEF5EF"
INPUT_FILL = "FFF8DC"

_thin = Side(style="thin", color="7F7F7F")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def f(size: int = 10, bold: bool = False, color: str = "000000", italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def cell(ws, ref, value=None, font=None, align=None, border=None, bg=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font or f()
    if align:
        c.alignment = align
    if border:
        c.border = border
    if bg:
        c.fill = fill(bg)
    return c


def boxed_range(ws, first_col: int, last_col: int, row: int, bg: str | None = None) -> None:
    for col in range(first_col, last_col + 1):
        c = ws.cell(row=row, column=col)
        c.border = BOX
        if bg:
            c.fill = fill(bg)


def merge(ws, c1: int, r1: int, c2: int, r2: int) -> None:
    if (c1, r1) != (c2, r2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def print_setup(ws: Worksheet) -> None:
    """One landscape Letter page per sheet."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    for side in ("left", "right", "top", "bottom"):
        setattr(ws.page_margins, side, 0.4)
    ws.print_options.horizontalCentered = True


def absent_rows(ws: Worksheet, rng: str, here_col: str, first_row: int, marker: str) -> None:
    """Grey out a whole row once its Here cell says the kid is away.

    The reference is written column-absolute but row-relative ($A15, not
    $A$15), so each row tests its own Here cell rather than the first one's.
    """
    ws.conditional_formatting.add(
        rng,
        FormulaRule(
            formula=[f'UPPER(${here_col}{first_row})="{marker}"'],
            fill=fill(GREY),
            stopIfTrue=False,
        ),
    )


def fair_colors(ws: Worksheet, rng: str) -> None:
    """Shade kids well under an even split (owed time) or well over it."""
    ws.conditional_formatting.add(
        rng, CellIsRule(operator="lessThanOrEqual", formula=["-3"], fill=fill(PALE_RED))
    )
    ws.conditional_formatting.add(
        rng, CellIsRule(operator="greaterThanOrEqual", formula=["3"], fill=fill(PALE_GREEN))
    )
