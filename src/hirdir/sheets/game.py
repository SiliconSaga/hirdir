"""One sheet per game: a practice plan on top, minute-by-minute play below.

The coach writes stopwatch minutes into In/Out pairs during the game. Minutes
played and ± fair are formulas, so typing the sheet up afterwards is enough to
feed the Season tab.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from ..config import Game, TeamConfig
from ..layout import (
    FIRST_PLAYER_ROW,
    SETTINGS_ROW,
    GameColumns,
    SettingsField,
    player_rows,
    settings_cells,
)
from ..style import (
    BOX,
    CENTER,
    GREEN,
    INPUT_FILL,
    INPUT_TEXT,
    LEFT,
    LIGHT,
    LINK,
    RIGHT,
    absent_rows,
    boxed_range,
    cell,
    f,
    fair_colors,
    fill,
    merge,
    print_setup,
)

PRACTICE_ROW = 5
ABSENT = "A"  # written in the Here column, before or after the game
SETTINGS_LABELS = ["Game ended at minute:", "Players per side:", "Fair share per kid:"]
LEGEND = (
    "In / Out = stopwatch minute. Back on again? Use the next In/Out pair. Still on at the end? "
    "Leave Out blank and fill in “Game ended at minute” — it counts to the end. Here ✓ even if a "
    "kid refuses to go in, so they still count toward fair share.",
    "Here: ✓ came · A didn't come (grey row, left out of fair share — mark it ahead of time if you "
    "already know) · blank = not recorded. ± fair = minutes vs an even split (game length × "
    "players per side ÷ kids here); red = owed time, green = had extra. ★ = doing great, can take "
    "less attention · Shy / Needs help = give extra attention. The Season tab adds it all up.",
)


def _settings_row(ws, cfg: TeamConfig, fields: list[SettingsField], kids_count: str) -> None:
    r = SETTINGS_ROW
    ws.row_dimensions[r].height = 24
    end, side, share = fields
    values = [
        (end, None, f(13, True, INPUT_TEXT), INPUT_FILL, None),
        (side, cfg.on_field, f(13, True, INPUT_TEXT), INPUT_FILL, None),
        (
            share,
            f'=IF(OR({end.value_ref}="",{kids_count}=0),"",'
            f"ROUND({end.value_ref}*{side.value_ref}/{kids_count},1))",
            f(13, True),
            None,
            '0.0" min"',
        ),
    ]
    for field, value, font, bg, number_format in values:
        c1, c2 = field.label_span
        cell(ws, f"{L(c1)}{r}", field.label, f(10, True), RIGHT)
        merge(ws, c1, r, c2, r)
        v1, v2 = field.value_span
        c = cell(ws, f"{L(v1)}{r}", value, font, CENTER, BOX, bg)
        boxed_range(ws, v1, v2, r, bg)
        merge(ws, v1, r, v2, r)
        if number_format:
            c.number_format = number_format


def _practice_section(ws, cfg: TeamConfig, cols: GameColumns, activity_ref: str) -> int:
    r = PRACTICE_ROW
    cell(ws, f"A{r}", cfg.practice_title, f(11, True, "FFFFFF"), LEFT, None, GREEN)
    merge(ws, 1, r, cols.last, r)
    r += 1
    spans = [
        ("Order", 1, 1),
        ("Activity (pick from the Activities tab, or write your own)", 2, 6),
        ("Min", 7, 8),
        ("Kids liked it?\n+  ~  −", 9, 10),
        ("Useful?\n+  ~  −", 11, 12),
        ("Notes", 13, cols.last),
    ]
    for label, c1, c2 in spans:
        cell(ws, f"{L(c1)}{r}", label, f(9, True), CENTER, None, LIGHT)
        boxed_range(ws, c1, c2, r, LIGHT)
        merge(ws, c1, r, c2, r)
    ws.row_dimensions[r].height = 26

    dv = DataValidation(
        type="list", formula1=activity_ref, allow_blank=True, showErrorMessage=False
    )
    ws.add_data_validation(dv)
    for i in range(cfg.practice_rows):
        r += 1
        for _label, c1, c2 in spans:
            boxed_range(ws, c1, c2, r)
            merge(ws, c1, r, c2, r)
            ws.cell(row=r, column=c1).font = f(11)
            ws.cell(row=r, column=c1).alignment = CENTER if c1 == 1 else LEFT
        ws.cell(row=r, column=1).value = i + 1
        dv.add(f"B{r}")
        ws.row_dimensions[r].height = 22
    return r


def _column_widths(cols: GameColumns) -> dict[int, float]:
    widths = {
        cols.here: 6,
        cols.num: 5,
        cols.name: 16,
        cols.minutes: 8,
        cols.fair: 6,
        cols.goals: 9,
        cols.star: 5,
        cols.shy: 5,
        cols.help: 7.5,  # enough that "Needs help" doesn't break mid-word
        cols.notes: 30,
    }
    for ci, co in cols.stints:
        widths[ci] = widths[co] = 5.5
    return widths


def _minutes_formula(cols: GameColumns, row: int, end_ref: str) -> str:
    """Sum of (Out − In) per stint; a blank Out means 'still on at the end'.

    Blank unless the kid was there: no Here mark and no stamps, or an explicit
    "A" for absent. Blank keeps them out of the fair-share divisor, so an absent
    kid is not recorded as owed time. A kid marked absent who turns up anyway
    still gets minutes once stamps are written.
    """
    here = f"{L(cols.here)}{row}"
    stamps = f"{L(cols.first_stint_col)}{row}:{L(cols.last_stint_col)}{row}"
    parts = "+".join(
        f'IF({L(ci)}{row}="",0,MAX(0,IF({L(co)}{row}="",{end_ref},{L(co)}{row})-{L(ci)}{row}))'
        for ci, co in cols.stints
    )
    away = f'OR({here}="",UPPER({here})="{ABSENT}")'
    return f'=IF(AND({away},COUNT({stamps})=0),"",{parts})'


def build(wb, cfg: TeamConfig, index: int, game: Game, cols: GameColumns, activity_ref: str):
    ws = wb.create_sheet(cfg.sheet_names[index])
    last = cols.last

    where = f" · Field {game.field}" if game.field else ""
    cell(
        ws,
        "A1",
        f"{cfg.team} · Game {index + 1} of {len(cfg.games)} · "
        f"{game.date:%a %b %-d} vs {game.opponent}{where}",
        f(16, True, GREEN),
        LEFT,
    )
    merge(ws, 1, 1, last, 1)
    ws.row_dimensions[1].height = 26
    cell(
        ws,
        "A2",
        "Coach: ______________    Other coach: ______________    "
        "Start a stopwatch at kickoff — pause it for the break.",
        f(10),
        LEFT,
    )
    merge(ws, 1, 2, last, 2)

    widths = _column_widths(cols)
    end, side, share = settings_cells(widths, SETTINGS_LABELS)

    rows = player_rows(cfg)
    grid_top = _practice_section(ws, cfg, cols, activity_ref) + 2
    cell(
        ws,
        f"A{grid_top}",
        "GAME — write the stopwatch minute when a kid goes In and comes Out",
        f(11, True, "FFFFFF"),
        LEFT,
        None,
        GREEN,
    )
    merge(ws, 1, grid_top, last, grid_top)

    head_row = grid_top + 1
    heads = {
        cols.here: "Here\n✓ / A",
        cols.num: "#",
        cols.name: "Player",
        cols.minutes: "Minutes\nplayed",
        cols.fair: "±\nfair",
        cols.goals: "Goals\n(tally)",
        cols.star: "★",
        cols.shy: "Shy",
        cols.help: "Needs\nhelp",
        cols.notes: "Notes\n(refused? quit early?)",
    }
    for i, (ci, co) in enumerate(cols.stints, 1):
        heads[ci] = f"In\n{i}"
        heads[co] = f"Out\n{i}"
    for col, label in heads.items():
        cell(ws, f"{L(col)}{head_row}", label, f(9, True), CENTER, BOX, LIGHT)
    ws.row_dimensions[head_row].height = 30

    first_row = head_row + 1
    for i in range(rows):
        row = first_row + i
        roster_row = FIRST_PLAYER_ROW + i
        for col in range(1, last + 1):
            c = ws.cell(row=row, column=col)
            c.border = BOX
            c.alignment = CENTER
            c.font = f(12)
        for pair, (ci, co) in enumerate(cols.stints):
            if pair % 2:  # shade alternate In/Out pairs so the eye can track them
                ws.cell(row=row, column=ci).fill = fill(LIGHT)
                ws.cell(row=row, column=co).fill = fill(LIGHT)

        num = ws.cell(row=row, column=cols.num)
        num.value = f'=IF(Roster!$A${roster_row}="","",Roster!$A${roster_row})'
        num.font = f(13, True, LINK)
        name = ws.cell(row=row, column=cols.name)
        name.value = f'=IF(Roster!$B${roster_row}="","",Roster!$B${roster_row})'
        name.font = f(13, True, LINK)
        name.alignment = LEFT

        minutes = ws.cell(row=row, column=cols.minutes)
        minutes.value = _minutes_formula(cols, row, end.value_ref)
        minutes.font = f(13, True)
        minutes.number_format = "0"
        fair = ws.cell(row=row, column=cols.fair)
        fair.value = (
            f'=IF(OR({L(cols.minutes)}{row}="",{share.value_ref}=""),"",'
            f"ROUND({L(cols.minutes)}{row}-{share.value_ref},0))"
        )
        fair.number_format = "+0;-0;0"
        fair.font = f(11)
        ws.cell(row=row, column=cols.notes).alignment = LEFT
        ws.cell(row=row, column=cols.notes).font = f(10)
        ws.row_dimensions[row].height = 26
    last_row = first_row + rows - 1

    kids = f'COUNTIF({L(cols.minutes)}{first_row}:{L(cols.minutes)}{last_row},">=0")'
    _settings_row(ws, cfg, [end, side, share], kids)
    fair_colors(ws, f"{L(cols.fair)}{first_row}:{L(cols.fair)}{last_row}")
    absent_rows(
        ws,
        f"{L(cols.here)}{first_row}:{L(last)}{last_row}",
        L(cols.here),
        first_row,
        ABSENT,
    )

    row = last_row + 2
    for line in LEGEND:
        cell(ws, f"A{row}", line, f(9, italic=True), LEFT)
        merge(ws, 1, row, last, row)
        ws.row_dimensions[row].height = 26
        row += 1

    for col, width in widths.items():
        ws.column_dimensions[L(col)].width = width
    print_setup(ws)
    return first_row
