from datetime import date

import pytest
from openpyxl import load_workbook

from hirdir import config, workbook
from hirdir.layout import FIRST_PLAYER_ROW, SPARE_ROWS, game_columns, settings_cells
from hirdir.sheets.game import SETTINGS_LABELS, _column_widths


def test_tabs_are_roster_season_activities_then_games(cfg, workbook_file):
    wb = load_workbook(workbook_file)
    assert wb.sheetnames == ["Roster", "Season", "Activities", *cfg.sheet_names]


def test_roster_holds_no_birthdates_or_ages(cfg, workbook_file):
    """Birthdates order the lineup and are then dropped — they never reach the file."""
    wb = load_workbook(workbook_file)
    text = " ".join(
        str(c.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    ).lower()
    assert "birth" not in text
    assert "years old" not in text
    for player in cfg.players:
        assert player.dob is not None  # the config knew them...
        assert str(player.dob) not in text  # ...the workbook does not
        assert str(player.dob.year) not in text


def test_roster_lists_players_youngest_first_with_spare_rows(cfg, workbook_file):
    ws = load_workbook(workbook_file)["Roster"]
    names = [ws.cell(row=FIRST_PLAYER_ROW + i, column=2).value for i in range(len(cfg.players))]
    assert names == [p.name for p in cfg.players]
    for spare in range(SPARE_ROWS):
        row = FIRST_PLAYER_ROW + len(cfg.players) + spare
        assert ws.cell(row=row, column=2).value is None
        assert ws.cell(row=row, column=2).border.left.style == "thin"


def test_game_sheet_titles_name_the_opponent_and_field(cfg, workbook_file):
    wb = load_workbook(workbook_file)
    title = wb[cfg.sheet_names[0]]["A1"].value
    assert "vs Otters" in title
    assert "Field 1" in title
    assert "Game 1 of 3" in title


def test_game_rows_pull_names_from_the_roster(cfg, workbook_file):
    cols = game_columns(cfg)
    ws = load_workbook(workbook_file)[cfg.sheet_names[0]]
    first = _first_player_row(ws, cols)
    for i in range(len(cfg.players) + SPARE_ROWS):
        formula = ws.cell(row=first + i, column=cols.name).value
        assert f"Roster!$B${FIRST_PLAYER_ROW + i}" in formula


def test_minutes_sum_every_stint_and_run_to_the_end_minute(cfg, workbook_file):
    cols = game_columns(cfg)
    ws = load_workbook(workbook_file)[cfg.sheet_names[0]]
    row = _first_player_row(ws, cols)
    formula = ws.cell(row=row, column=cols.minutes).value
    assert formula.count("MAX(0,") == cfg.stints  # one term per In/Out pair
    assert "$D$3" in formula  # blank Out counts to the game's end minute


def test_fair_share_divides_the_game_among_the_kids_who_came(cfg, workbook_file):
    cols = game_columns(cfg)
    ws = load_workbook(workbook_file)[cfg.sheet_names[0]]
    end, side, share = settings_cells(_column_widths(cols), SETTINGS_LABELS)
    assert ws[side.value_ref.replace("$", "")].value == cfg.on_field
    formula = ws[share.value_ref.replace("$", "")].value
    assert f"{end.value_ref}*{side.value_ref}/COUNTIF" in formula


@pytest.mark.parametrize("stints", [2, 4, 6])
def test_settings_labels_have_room_for_their_text(example_data, tmp_path, stints):
    """A label merged too narrow spills over its neighbour — at any stint count."""
    example_data["stints"] = stints
    cfg = config.parse(example_data)
    path = workbook.write(cfg, tmp_path / f"s{stints}.xlsx")
    ws = load_workbook(path)[cfg.sheet_names[0]]
    merged = {str(r) for r in ws.merged_cells.ranges}
    widths = _column_widths(game_columns(cfg))
    for field in settings_cells(widths, SETTINGS_LABELS):
        c1, c2 = field.label_span
        assert f"{_letter(c1)}3:{_letter(c2)}3" in merged
        span = sum(widths[c] for c in range(c1, c2 + 1))
        assert span >= len(field.label), f"{field.label!r} needs more room"
        # the label must not run into the value cell next to it
        assert field.value_span[0] > c2


def test_season_reads_every_game_sheet(cfg, workbook_file):
    ws = load_workbook(workbook_file)["Season"]
    row = FIRST_PLAYER_ROW
    formulas = " ".join(str(c.value) for c in ws[row])
    for name in cfg.sheet_names:
        assert f"'{name}'!" in formulas


def test_sheets_print_on_one_landscape_page(workbook_file):
    for ws in load_workbook(workbook_file).worksheets:
        assert ws.page_setup.orientation == "landscape"
        assert ws.page_setup.fitToWidth == 1
        assert ws.page_setup.fitToHeight == 1


def test_practice_rows_offer_the_activities_list(cfg, workbook_file):
    ws = load_workbook(workbook_file)[cfg.sheet_names[0]]
    validations = ws.data_validations.dataValidation
    assert len(validations) == 1
    assert validations[0].formula1.startswith("Activities!$A$")
    covered = str(validations[0].sqref).split()
    assert len(covered) == cfg.practice_rows
    assert all(ref.startswith("B") for ref in covered)


def test_absent_rows_grey_out_per_row_not_all_at_once(cfg, workbook_file):
    """The rule must be row-relative, or one kid marked A greys the whole team."""
    cols = game_columns(cfg)
    ws = load_workbook(workbook_file)[cfg.sheet_names[0]]
    first = _first_player_row(ws, cols)
    rules = [
        (rng, rule)
        for rng in ws.conditional_formatting
        for rule in rng.rules
        if rule.type == "expression"
    ]
    assert len(rules) == 1
    rng, rule = rules[0]
    assert rule.formula == [f'UPPER($A{first})="A"']  # $A15, never $A$15
    assert str(rng.sqref) == f"A{first}:{_letter(cols.last)}{first + len(cfg.players) + 1}"


def test_written_file_lands_where_asked(cfg, tmp_path):
    out = workbook.write(cfg, tmp_path / "nested" / "book.xlsx")
    assert out.exists()


def test_games_without_a_field_still_build(example_data, tmp_path):
    from hirdir import config

    example_data["games"] = [{"date": "2026-09-20", "opponent": "Otters"}]
    cfg = config.parse(example_data)
    ws = load_workbook(workbook.write(cfg, tmp_path / "b.xlsx"))[cfg.sheet_names[0]]
    assert "Field" not in ws["A1"].value
    assert cfg.games[0].date == date(2026, 9, 20)


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)


def _first_player_row(ws, cols) -> int:
    """The row under the grid header — found the way a reader would."""
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=cols.name).value == "Player":
            return row + 1
    raise AssertionError("game grid header not found")
