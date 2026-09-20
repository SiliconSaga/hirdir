"""Evaluate a filled-in workbook for real, the way a spreadsheet app would.

Uses the `formulas` package (dev dependency), so a wrong cell reference or a
function our tooling can't evaluate fails here rather than at the field.
"""

from __future__ import annotations

import re

import pytest
from openpyxl import load_workbook

from hirdir.bake import bake
from hirdir.layout import FIRST_PLAYER_ROW, game_columns, player_rows, settings_cells
from hirdir.sheets.game import SETTINGS_LABELS, _column_widths

formulas = pytest.importorskip("formulas")

CELL_KEY = re.compile(r"'\[[^\]]+\](.+)'!([A-Z]+\d+)$")
# Excel error values, as opposed to a cell whose text merely starts with "#"
# (the jersey-number column header is a literal "#").
EXCEL_ERRORS = {"#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NULL!", "#NUM!"}


def fair_share_cell(cfg) -> str:
    """Where "Fair share per kid" lands — it moves with the column widths."""
    fields = settings_cells(_column_widths(game_columns(cfg)), SETTINGS_LABELS)
    return fields[2].value_ref.replace("$", "")


def end_minute_cell(cfg) -> str:
    fields = settings_cells(_column_widths(game_columns(cfg)), SETTINGS_LABELS)
    return fields[0].value_ref.replace("$", "")


def first_grid_row(ws, cols) -> int:
    """The first player row of the game grid, found by its header."""
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=cols.name).value == "Player":
            return row + 1
    raise AssertionError("game grid header not found")


def evaluate(path) -> dict[tuple[str, str], object]:
    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    values = {}
    for key, value in solution.items():
        match = CELL_KEY.match(key)
        if match:
            cell = value.value[0, 0] if hasattr(value, "value") else value
            values[(match.group(1).upper(), match.group(2))] = cell
    return values


@pytest.fixture
def played(cfg, workbook_file):
    """A game as it actually goes: someone still on at the whistle, someone who refuses."""
    cols = game_columns(cfg)
    wb = load_workbook(workbook_file)
    g1, g2 = wb[cfg.sheet_names[0]], wb[cfg.sheet_names[1]]
    first = first_grid_row(g1, cols)

    def col(ws, row, column, value):
        ws.cell(row=row, column=column).value = value

    wb["Roster"]["A5"] = 1  # Dev (youngest) gets jersey 1
    g1[end_minute_cell(cfg)] = 30  # game ended at minute 30, 4 per side
    # Dev: on 0-8, back on at 15 and still on at the whistle = 23 minutes
    col(g1, first, cols.here, "x")
    col(g1, first, cols.stints[0][0], 0)
    col(g1, first, cols.stints[0][1], 8)
    col(g1, first, cols.stints[1][0], 15)
    col(g1, first, cols.goals, 1)
    col(g1, first, cols.shy, "x")
    # Gita: came, refused to go in
    col(g1, first + 1, cols.here, "x")
    # Bjorn: 0-10 and 20-30 = 20 minutes, coach forgot the Here tick
    col(g1, first + 2, cols.stints[0][0], 0)
    col(g1, first + 2, cols.stints[0][1], 10)
    col(g1, first + 2, cols.stints[1][0], 20)
    col(g1, first + 2, cols.stints[1][1], 30)
    # Cleo: 5-12 = 7 minutes
    col(g1, first + 3, cols.stints[0][0], 5)
    col(g1, first + 3, cols.stints[0][1], 12)
    # Hugo: known in advance not to be coming
    col(g1, first + 4, cols.here, "A")
    # Game 2: Dev turned up, nothing recorded yet
    col(g2, first, cols.here, "✓")
    wb.save(workbook_file)
    return workbook_file, cols, first


def test_no_formula_errors_in_a_fresh_workbook(cfg, workbook_file):
    errors = [
        (key, value)
        for key, value in evaluate(workbook_file).items()
        if isinstance(value, str) and value in EXCEL_ERRORS
    ]
    assert errors == []


def test_minutes_and_fair_share(cfg, played):
    path, cols, first = played
    values = evaluate(path)
    game = cfg.sheet_names[0].upper()

    def at(row_offset, column):
        return values[(game, f"{_letter(column)}{first + row_offset}")]

    assert at(0, cols.minutes) == 23  # still on at the whistle
    assert at(1, cols.minutes) == 0  # came but refused
    assert at(2, cols.minutes) == 20
    assert at(3, cols.minutes) == 7
    assert at(4, cols.minutes) == ""  # marked absent: stays blank
    assert at(5, cols.minutes) == ""  # nothing recorded: stays blank
    # 30 minutes × 4 per side ÷ 4 kids present
    assert values[(game, fair_share_cell(cfg))] == 30
    assert at(0, cols.fair) == -7
    assert at(1, cols.fair) == -30
    assert at(4, cols.fair) == ""  # absent kids aren't owed time


def test_season_rolls_up_minutes_goals_and_flags(cfg, played):
    path, _cols, _first = played
    values = evaluate(path)
    season = "SEASON"
    row = FIRST_PLAYER_ROW  # Dev
    assert values[(season, f"A{row}")] == 1  # jersey pulled from the roster
    assert values[(season, f"B{row}")] == "Dev"
    assert values[(season, f"C{row}")] == 23  # game 1 minutes
    assert values[(season, f"D{row}")] == 0  # game 2: here, no minutes yet
    totals = 3 + len(cfg.games)
    assert values[(season, f"{_letter(totals)}{row}")] == 2  # games
    assert values[(season, f"{_letter(totals + 1)}{row}")] == 0  # missed
    assert values[(season, f"{_letter(totals + 2)}{row}")] == 23  # minutes
    assert values[(season, f"{_letter(totals + 3)}{row}")] == 11.5  # per game
    assert values[(season, f"{_letter(totals + 4)}{row}")] == -7  # ± fair
    assert values[(season, f"{_letter(totals + 5)}{row}")] == 1  # goals
    assert values[(season, f"{_letter(totals + 7)}{row}")] == 1  # shy


def test_an_absent_kid_shows_as_A_and_counts_as_missed(cfg, played):
    path, _cols, _first = played
    values = evaluate(path)
    season = "SEASON"
    row = FIRST_PLAYER_ROW + 4  # Hugo, marked A for game 1
    assert values[(season, f"B{row}")] == "Hugo"
    assert values[(season, f"C{row}")] == "A"
    totals = 3 + len(cfg.games)
    assert values[(season, f"{_letter(totals)}{row}")] == 0  # not counted as a game played
    assert values[(season, f"{_letter(totals + 1)}{row}")] == 1  # missed one
    # fair share is still 30: four kids were there, Hugo wasn't
    assert values[(cfg.sheet_names[0].upper(), fair_share_cell(cfg))] == 30
    footer = FIRST_PLAYER_ROW + player_rows(cfg)
    assert values[(season, f"C{footer}")] == 4  # kids at game 1


def test_bake_writes_values_that_a_non_calculating_viewer_can_read(cfg, played):
    path, cols, first = played
    written = bake(path)
    assert written > 0
    sheet = cfg.sheet_names[0]
    cached = load_workbook(path, data_only=True)[sheet]
    assert cached.cell(row=first, column=cols.minutes).value == 23
    live = load_workbook(path)[sheet]
    assert live.cell(row=first, column=cols.minutes).value.startswith("=")


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)
